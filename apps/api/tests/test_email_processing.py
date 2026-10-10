import io
from datetime import datetime, timezone

from docx import Document

from app.services.email_processing import (
    build_ai_input,
    extract_attachment_text,
    process_email_content,
)


def test_email_content_is_sanitized_and_normalized_for_ai() -> None:
    processed = process_email_content(
        body_html=(
            '<script>alert("x")</script><p style="color: #333; position: fixed">'
            "Meeting: <strong>12 Oct</strong></p>"
            '<img src="https://tracker.example/open.gif" width="1" height="1">'
            '<img src="https://cdn.example/hero.jpg" width="640" height="320">'
            '<p>Register at <a href="https://example.com">this link</a>.</p>'
        ),
        body_text=(
            "Meeting: 12 Oct\nRegister at https://example.com\n\n"
            "On Monday, Alice wrote:\n> old repeated message"
        ),
        sent_at=datetime(2026, 10, 10, 16, 30, tzinfo=timezone.utc),
        user_timezone="Asia/Singapore",
    )

    assert "script" not in (processed.safe_html or "")
    assert "tracker.example" not in (processed.safe_html or "")
    assert "https://cdn.example/hero.jpg" in (processed.safe_html or "")
    assert 'href="https://example.com"' in (processed.safe_html or "")
    assert "color: #333" in (processed.safe_html or "")
    assert "position" not in (processed.safe_html or "")
    assert "12 Oct" in processed.cleaned_text
    assert "https://example.com" in processed.cleaned_text
    assert "old repeated message" not in processed.cleaned_text
    assert processed.removed_sections == ["quoted_history"]
    assert processed.reference_date == "2026-10-11"
    assert processed.original_timezone == "UTC"
    assert processed.user_timezone == "Asia/Singapore"


def test_invalid_user_timezone_falls_back_to_utc() -> None:
    processed = process_email_content(
        body_html=None,
        body_text="Reply tomorrow.",
        sent_at=datetime(2026, 10, 10, 23, 0, tzinfo=timezone.utc),
        user_timezone="Not/AZone",
    )

    assert processed.user_timezone == "UTC"
    assert processed.reference_date == "2026-10-10"


def test_html_only_email_keeps_link_target_in_ai_text() -> None:
    processed = process_email_content(
        body_html='<p>Register using <a href="https://example.com/form">this form</a>.</p>',
        body_text=None,
        sent_at=datetime(2026, 10, 10, tzinfo=timezone.utc),
        user_timezone="UTC",
    )

    assert "https://example.com/form" in processed.cleaned_text


def test_txt_attachment_extraction_and_size_limit() -> None:
    extracted = extract_attachment_text(
        content="Amount: S$15\nDate: 12 October".encode(),
        filename="details.txt",
        mime_type="text/plain",
        max_size_bytes=1024,
    )
    rejected = extract_attachment_text(
        content=b"too large",
        filename="details.txt",
        mime_type="text/plain",
        max_size_bytes=2,
    )

    assert extracted.status == "extracted"
    assert extracted.extracted_text == "Amount: S$15\nDate: 12 October"
    assert rejected.status == "too_large"


def test_docx_attachment_extracts_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Project deadline: 18 October")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Mengyi"
    stream = io.BytesIO()
    document.save(stream)

    extracted = extract_attachment_text(
        content=stream.getvalue(),
        filename="brief.docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        max_size_bytes=1024 * 1024,
    )

    assert extracted.status == "extracted"
    assert "Project deadline: 18 October" in (extracted.extracted_text or "")
    assert "Owner Mengyi" in (extracted.extracted_text or "")


def test_unsupported_attachment_is_not_parsed() -> None:
    extracted = extract_attachment_text(
        content=b"MZ",
        filename="program.exe",
        mime_type="application/x-msdownload",
        max_size_bytes=1024,
    )

    assert extracted.status == "unsupported"
    assert extracted.error == "attachment_type_not_allowed"


def test_attachment_content_must_match_claimed_type() -> None:
    extracted = extract_attachment_text(
        content=b"not actually a PDF",
        filename="spoofed.pdf",
        mime_type="application/pdf",
        max_size_bytes=1024,
    )

    assert extracted.status == "failed"
    assert extracted.error == "attachment_content_mismatch"


def test_ai_input_only_contains_successfully_extracted_attachments() -> None:
    payload = build_ai_input(
        subject="Meeting",
        sender={"name": "Alice", "email": "alice@example.com"},
        recipients={"to": [{"name": "Bob", "email": "bob@example.com"}]},
        cleaned_text="Meet tomorrow.",
        sent_at=datetime(2026, 10, 10, tzinfo=timezone.utc),
        reference_date="2026-10-10",
        user_timezone="Asia/Singapore",
        attachments=[
            {
                "filename": "agenda.txt",
                "mime_type": "text/plain",
                "extraction_status": "extracted",
                "extracted_text": "Agenda item one",
            },
            {
                "filename": "program.exe",
                "mime_type": "application/x-msdownload",
                "extraction_status": "unsupported",
                "extracted_text": None,
            },
        ],
    )

    assert len(payload["attachments"]) == 1  # type: ignore[arg-type]
    assert "raw_body_html" not in payload
